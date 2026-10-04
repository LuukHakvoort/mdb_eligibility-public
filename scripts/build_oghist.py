"""Fetch and parse the World Bank's historical country classification file
into data/extracted/oghist_income.csv and
data/extracted/oghist_lending_category.csv -- Component D's IBRD/IDA source
(docs/architecture.md section 2; docs/manual_data_checklist.md Phase 3).

**Rewritten 2026-09-27** (docs/coding_decisions.md) to close a gap this
project had documented as unfillable: calendar years 2023-2025 had neither an
income group (OGHIST.xlsx ended at FY2023 -> 2022) nor a lending category
(the "current classification" file was a single present-day snapshot,
FY27 -> 2026 only, not projected backward). Both turned out to already be one
click away, not missing from the World Bank at all -- the two files this
script used to fetch were each a stale or partial view of the same
underlying data the Bank now publishes as one tidy panel:

**`CLASS_hist` ("Historical classification by income, lending, and fragile
status")**, `ddh-openapi.worldbank.org/resources/DR0095347/download` (found
via the WDI dataset's own catalog page, `datacatalog.worldbank.org/search/
dataset/0037712` -- the OGHIST.xlsx URL this script used to hit,
`databankfiles.worldbank.org/public/ddpext_download/site-content/
OGHIST.xlsx`, is simply not the current file the Bank publishes there
anymore; the DDH catalog is). Its `list` sheet is already tidy: one row per
(country, calendar `year`), with `income` (a plain string: "Low income" /
"Lower middle income" / "Upper middle income" / "High income") and `lending`
(a direct `IDA`/`IBRD`/`Blend` column, no group-vs-category ambiguity)
side by side, **for every year 1988-2026** including 2023-2025. **Verified,
not assumed**, before trusting it over the two old sources: its `income`
column matches the old `oghist_income.csv` (built from the stale OGHIST.xlsx)
on every one of 4,649 overlapping (iso3, year) pairs, zero mismatches; its
`lending` column is populated from 2002 onward -- this project's entire study
window -- not just the current year, so the old "current snapshot only, not
projected backward" limitation is gone, and with it most of the reason
`components/window.py`'s income-group hedge (blend ambiguity, small-state
exception) was ever needed for years inside 2002-2026: a direct lending
category is now available for nearly every classified country in every
study year, not just 2026. `components/window.py`'s hedge logic is
unchanged (still correct, still needed for whatever residual gap remains --
e.g. a country genuinely absent from `lending` some year), just exercised far
less often now.

This script therefore fetches one file, not two, and writes both output CSVs
from it:

- `oghist_income.csv`: `income` mapped to the same L/LM/UM/H codes the old
  OGHIST.xlsx sheet used (`_GROUP_MAP` below), for every (iso3, year) inside
  MIN_YEAR-1..MAX_YEAR (the extra year below MIN_YEAR matches this project's
  existing convention -- see the old `_parse_country_analytical_history`,
  now retired).
- `oghist_lending_category.csv`: `lending` directly, same year range, with a
  `note` recording this is the historical (not current-snapshot) series --
  the field `components/window.py`'s `_oghist_window_ok` already reads first,
  before falling back to income group.

Usage:  python3 scripts/build_oghist.py
"""

from __future__ import annotations

import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _panel import EXTRACTED_DIR, MAX_YEAR, MIN_YEAR, REPO, REPORTS_DIR, write_csv  # noqa: E402
from _reference import ISO3_CODES  # noqa: E402

RAW_DIR = REPO / "data" / "raw" / "oghist"
CLASS_HIST_URL = "https://ddh-openapi.worldbank.org/resources/DR0095347/download"
USER_AGENT = "mdb-eligibility-thesis-research/1.0 (+academic research; contact ljr.hakvoort@gmail.com)"
RETRIES = 4
BACKOFF_BASE = 2.0

INCOME_CSV = EXTRACTED_DIR / "oghist_income.csv"
LENDING_CATEGORY_CSV = EXTRACTED_DIR / "oghist_lending_category.csv"
UNRESOLVED_REPORT = REPORTS_DIR / "oghist_unresolved.csv"

_GROUP_MAP = {
    "Low income": "L", "Lower middle income": "LM",
    "Upper middle income": "UM", "High income": "H",
}
_LENDING_VALUES = {"IDA", "IBRD", "Blend"}


def _fetch(url: str, cache_name: str) -> Path:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    cache_path = RAW_DIR / cache_name
    if cache_path.is_file():
        return cache_path

    last_error: Exception | None = None
    for attempt in range(RETRIES):
        try:
            request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(request, timeout=60) as response:
                body = response.read()
            tmp = cache_path.with_suffix(".tmp")
            tmp.write_bytes(body)
            tmp.rename(cache_path)
            return cache_path
        except (urllib.error.URLError, TimeoutError) as exc:
            last_error = exc
            if attempt < RETRIES - 1:
                time.sleep(BACKOFF_BASE * (2**attempt))
    raise SystemExit(f"error: could not fetch {url} after {RETRIES} attempts ({last_error})")


def _parse_class_hist(path: Path) -> tuple[list[dict], list[dict], list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["list"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}

    income_rows: list[dict] = []
    lending_rows: list[dict] = []
    unresolved: list[dict] = []
    seen_unresolved: set[str] = set()
    for r in rows[1:]:
        code = r[idx["country"]]
        year = r[idx["year"]]
        if not isinstance(code, str) or not code.strip() or not isinstance(year, int):
            continue
        code = code.strip()
        if year < MIN_YEAR - 1 or year > MAX_YEAR:
            continue
        if code not in ISO3_CODES:
            if code not in seen_unresolved:
                seen_unresolved.add(code)
                unresolved.append({
                    "iso3_attempted": code, "name": r[idx["country_name"]] or "",
                    "reason": "not a recognized ISO-3 code (likely a dissolved/historical/aggregate entity)",
                })
            continue

        income = r[idx["income"]]
        if isinstance(income, str) and income in _GROUP_MAP:
            income_rows.append({
                "iso3": code, "calendar_year": year, "income_group": _GROUP_MAP[income],
                "source": CLASS_HIST_URL,
            })

        lending = r[idx["lending"]]
        if isinstance(lending, str) and lending in _LENDING_VALUES:
            lending_rows.append({
                "iso3": code, "calendar_year": year, "lending_category": lending,
                "source": CLASS_HIST_URL,
                "note": "Historical series (2002 on), not a single current-year snapshot.",
            })
    return income_rows, lending_rows, unresolved


def main() -> int:
    path = _fetch(CLASS_HIST_URL, "CLASS_hist.xlsx")
    income_rows, lending_rows, unresolved = _parse_class_hist(path)

    write_csv(INCOME_CSV, ["iso3", "calendar_year", "income_group", "source"], income_rows)
    write_csv(UNRESOLVED_REPORT, ["iso3_attempted", "name", "reason"], unresolved)
    print(f"{len(income_rows)} income rows -> {INCOME_CSV.relative_to(REPO)} "
          f"({len(unresolved)} unresolved historical/dissolved/aggregate entities -> {UNRESOLVED_REPORT.relative_to(REPO)})")

    write_csv(LENDING_CATEGORY_CSV, ["iso3", "calendar_year", "lending_category", "source", "note"], lending_rows)
    years = sorted({row["calendar_year"] for row in lending_rows})
    print(f"{len(lending_rows)} lending-category rows ({years[0]}-{years[-1]}) -> {LENDING_CATEGORY_CSV.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
